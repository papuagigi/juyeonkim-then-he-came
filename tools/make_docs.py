# -*- coding: utf-8 -*-
"""
제출 문서 만들기 — 2026 금융 AI Challenge 양식(첨부1 기획서 · 첨부2 기능명세서)을 그대로 본뜬 Word 파일을 만들고
Word(COM)로 PDF 로 바꾼다.

실행: python tools/make_docs.py            → docs/ 아래 .docx 와 .pdf
필요: python-docx, Windows + Microsoft Word(PDF 변환)
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCS = os.path.join(ROOT, "docs")
IMG = os.path.join(DOCS, "img")
os.makedirs(DOCS, exist_ok=True)

TEAM = "김주연 등장"
MEMBERS = "김주연 (팀장, 1인 참가)"
URL = "https://papuagigi.github.io/juyeonkim-then-he-came/"
REPO = "https://github.com/papuagigi/juyeonkim-then-he-came"
FONT = "맑은 고딕"

meta = json.load(open(os.path.join(ROOT, "data", "meta.json"), encoding="utf-8"))
N_TERMS = len(re.findall(r"\{\s*t:\s*\"", open(os.path.join(ROOT, "assets", "terms.js"), encoding="utf-8").read()))
SAMPLES = {}
sp = os.path.join(DOCS, "sample_results.json")
if os.path.exists(sp):
    SAMPLES = json.load(open(sp, encoding="utf-8"))
AI_NOTE = os.environ.get("AI_SERVER_NOTE", "")


# ---------------- 서식 도우미 ----------------
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, sz=6, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def font(run, size=10, bold=False, color=None, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(cell_or_doc, text="", size=10, bold=False, color=None, align=None, space_after=2, indent=None, italic=False):
    p = cell_or_doc.add_paragraph()
    if text:
        r = p.add_run(text)
        font(r, size, bold, color, italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def rich(cell, parts, size=10, indent=None, space_after=2, bullet=False):
    """parts: [(text, bold), ...] 또는 문자열. **굵게** 표기를 지원한다."""
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bullet:
        p.paragraph_format.left_indent = Cm(0.5 if indent is None else indent)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        parts = "- " + parts if isinstance(parts, str) else [("- ", False)] + parts
    elif indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if isinstance(parts, str):
        segs = re.split(r"(\*\*[^*]+\*\*)", parts)
        parts = [(s[2:-2], True) if s.startswith("**") else (s, False) for s in segs if s]
    for text, bold in parts:
        r = p.add_run(text)
        font(r, size, bold)
    return p


def clear_cell(cell):
    for p in cell.paragraphs:
        p._element.getparent().remove(p._element)


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = FONT; st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); st.font.size = Pt(10)
    # 바닥글 쪽번호 "- 1 -"
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("- "); font(r, 9)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = fp.add_run(); font(run, 9)
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; run._r.append(el)
    r = fp.add_run(" -"); font(r, 9)
    return doc


def header_block(doc, attach_no, title):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c0, c1 = t.rows[0].cells
    c0.width = Cm(3.0); c1.width = Cm(13.6)
    for c in (c0, c1):
        set_cell_borders(c, sz=12); clear_cell(c)
    p = c0.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(f"첨부 {attach_no}"), 14, True)
    p = c1.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(title), 16, True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    # 팀명·구성원
    t2 = doc.add_table(rows=2, cols=2); t2.alignment = WD_TABLE_ALIGNMENT.CENTER; t2.autofit = False
    for i, (k, v) in enumerate((("팀명", TEAM), ("구성원 성명", MEMBERS))):
        a, b = t2.rows[i].cells
        a.width = Cm(4.6); b.width = Cm(12.0)
        for c in (a, b):
            set_cell_borders(c); clear_cell(c)
        set_cell_bg(a, "D9D9D9")
        p = a.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(k), 11, True)
        p = b.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(v), 11)
    para(doc, "( * 필수항목)", 9, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)


def section_table(doc):
    t = doc.add_table(rows=0, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    return t


def add_section(t, title, required=True):
    row = t.add_row(); c = row.cells[0]; c.width = Cm(16.6)
    set_cell_borders(c); set_cell_bg(c, "D9D9D9"); clear_cell(c)
    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    font(p.add_run(title), 11.5, True)
    if required:
        r = p.add_run("*"); font(r, 9, True); r.font.superscript = True
    row = t.add_row(); c = row.cells[0]; c.width = Cm(16.6)
    set_cell_borders(c); clear_cell(c)
    return c


def sub_table(cell, header, rows, widths, size=8.5):
    tbl = cell.add_table(rows=1, cols=len(header))
    tbl.autofit = False
    for i, h in enumerate(header):
        c = tbl.rows[0].cells[i]; c.width = Cm(widths[i]); set_cell_borders(c, 4, "808080"); set_cell_bg(c, "EDEDED"); clear_cell(c)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); font(p.add_run(h), size, True)
    for r in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(r):
            c = cells[i]; c.width = Cm(widths[i]); set_cell_borders(c, 4, "808080"); clear_cell(c)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            segs = re.split(r"(\*\*[^*]+\*\*)", str(v))
            for s in segs:
                if not s:
                    continue
                run = p.add_run(s[2:-2] if s.startswith("**") else s); font(run, size, s.startswith("**"))
    cell.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def images_row(cell, files, captions, width_cm):
    files = [f for f in files if os.path.exists(f)]
    if not files:
        return
    tbl = cell.add_table(rows=2, cols=len(files)); tbl.autofit = False
    for i, f in enumerate(files):
        c = tbl.rows[0].cells[i]; c.width = Cm(width_cm); clear_cell(c)
        p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(f, width=Cm(width_cm - 0.3))
        c2 = tbl.rows[1].cells[i]; c2.width = Cm(width_cm); clear_cell(c2)
        p = c2.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        font(p.add_run(captions[i]), 8, color="555555")


# ---------------- 첨부 1 기획서 ----------------
def build_plan():
    doc = new_doc()
    header_block(doc, 1, "2026 금융 AI Challenge 기획서")
    t = section_table(doc)

    c = add_section(t, "1. 서비스 명칭")
    rich(c, "**ETF 엑스레이(ETF X-ray)** — \"내 ETF 바구니, 겹치는 건 없을까?\"", 11)
    rich(c, "사회초년생이 산 ETF 여러 개를 '바구니'로 담으면 겹치는 종목·몰린 지역·특수 구조를 투시해 보여 주고, 생성형 AI가 계산된 근거만으로 쉬운 말 진단서를 써 주는 청년 자산형성 도우미 웹서비스")
    rich(c, f"웹서비스 URL: {URL}   ·   소스코드: {REPO}", 9)

    c = add_section(t, "2. 아이디어 기획 핵심내용(요약)")
    for s in [
        "**문제** — 2030 사회초년생은 ETF로 투자를 시작하지만, 이름이 다른 ETF를 여러 개 사서 '분산했다'고 믿는 경우가 많다. 실제로는 같은 지수·같은 종목(삼성전자·엔비디아)에 몰려 있는 **'가짜 분산'**이 흔하고, 상품명 속 낱말(커버드콜·TR·(H)·합성·레버리지)이 무슨 뜻인지 모른 채 산다.",
        "**해결** — 바구니를 담으면 ① 겹침 투시(ETF 쌍 겹침률, 여러 ETF에 함께 든 종목, 확인된 실효 비중) ② 구성 투시(투자 지역·자산 유형·위험등급 비율) ③ 발견 규칙 15종(같은 지수 중복, 70% 이상 겹침, 한 종목 집중, 레버리지·커버드콜·합성 구조, 소규모 순자산 등) ④ **생성형 AI 진단서**(근거만으로 쉬운 말, 숫자 사후 대조) ⑤ 이름 해독(사전 " + str(N_TERMS) + "낱말) ⑥ 같은 지수 다른 ETF 비교표를 한 화면에 보여 준다.",
        f"**데이터** — 국내 상장 ETF {meta['etfs']:,}종 마스터(기준일 2026-08-22) + ETF 구성종목 {meta['holding_pairs']:,}건(한국거래소에서 직접 수집, 기준일 2026-08-21) + 직접 작성한 이름 해독 사전.",
        "**원칙** — 추천·전망을 하지 않는다(사실 설명 + '확인해 볼 질문'). AI는 근거에 있는 숫자만 쓰고, 근거에 없는 숫자가 든 문장은 자동 삭제한다. 값이 없으면 없다고 표시한다. 개인정보는 수집하지 않는다.",
        "**MVP** — 설치·가입 없는 모바일 우선 웹. 예시 바구니 4종으로 30초 안에 체험할 수 있고, 바구니를 주소 하나로 공유한다. 화면 계산은 브라우저 안에서, AI 진단서는 서버(HyperCLOVA X)에서 처리하며 서버가 없어도 규칙 진단서로 항상 답한다.",
    ]:
        rich(c, s, bullet=True)

    c = add_section(t, "3. 문제 정의 및 제안 배경")
    rich(c, "**현재 금융 서비스의 구체적인 문제**", 10.5)
    for s in [
        f"**가짜 분산** — 국내 상장 ETF는 {meta['etfs']:,}종(2026-08 기준)이지만 코스피200(CR)을 따르는 ETF만 25종, 나스닥100 16종, S&P500 계열 24종이다. 브랜드(KODEX·TIGER·RISE·ACE)가 다르면 다른 상품으로 여기기 쉽다. 실제 계산으로 KODEX 200과 TIGER 200의 구성 겹침은 **99.8%**이고, 국내 대표지수 ETF 3종에 반도체 ETF를 더한 바구니는 삼성전자 한 종목의 실효 비중이 **30.8%**에 이른다. 나눠 샀지만 나뉘지 않았다.",
        "**이름을 모른 채 산다** — '월배당', '+7%프리미엄', '타겟커버드콜', 'TR', '(H)', '합성', '레버리지' 같은 낱말이 상품의 위험·수익 구조를 결정하지만, 매수 화면에는 설명이 없다. 분배수익률 14%를 '수익률 14%'로 오해해 사는 일이 흔하다.",
        "**물어볼 곳이 없다** — 소액·비대면 청년 투자자는 PB 상담 대상이 아니고, 증권사 앱은 '상품을 파는 화면'에 최적화되어 '내가 이미 산 것들이 서로 어떤 관계인지'는 보여 주지 않는다. 커뮤니티의 답은 근거가 없다.",
    ]:
        rich(c, s, bullet=True)
    rich(c, "**고객과 채널을 선택한 배경**", 10.5)
    for s in [
        "**고객: 2030 사회초년생·청년 투자자** — 첫 투자를 ETF로 시작하고, 연금저축·ISA 계좌에서 국내 상장 ETF를 사 모으는 사람들. 대회 세부 주제 '청년층 자산 형성을 위한 AI 기반 포용적 금융서비스'에 해당한다. 대상을 국내 상장 ETF로 정한 이유는 ① 연금저축·ISA에서는 국내 상장 ETF만 살 수 있어 청년의 보유가 여기에 몰리고 ② 구성종목이 매일 공시되어 '사실'로 계산할 수 있으며 ③ 해외 주식을 담은 국내 상장 ETF(TIGER 미국S&P500 등)까지 포함되어 청년 바구니의 대부분을 덮기 때문이다.",
        "**채널: 모바일 웹** — 설치·가입 없이 링크 하나로 열리고, 증권사 MTS의 '보유상품 진단' 메뉴나 카카오톡 공유 링크로 그대로 붙일 수 있다. 바구니가 주소(해시)에 담기므로 '내 바구니 봐줘'를 친구·가족·커뮤니티에 한 줄로 보낼 수 있다. 판단은 사용자가 하고 서비스는 투시만 하므로, 앱 안에 넣어도 부당권유 우려가 적다.",
    ]:
        rich(c, s, bullet=True)

    c = add_section(t, "4. 서비스 컨셉 및 차별성")
    rich(c, "**컨셉 — \"상품을 파는 화면이 아니라, 이미 산 것을 투시하는 화면\"** 엑스레이처럼 바구니 속을 있는 그대로 보여 주고, 판단은 사용자가 하도록 '확인해 볼 질문'으로 끝낸다.", 10)
    for s in [
        f"**구성종목 단위 투시** — 기존 앱(증권사 MTS, 핀테크 자산관리)은 상품 단위 수익률·비중만 보여 준다. 우리는 ETF 구성종목 {meta['holding_pairs']:,}건을 연결해 \"이름은 5개인데 속은 삼성전자 30.8%\"를 계산한다. 비중이 공개된 국내 주식은 **비중 기준(두 ETF의 종목별 비중 중 작은 값의 합)**, 비중이 없는 해외 주식은 **종목 수 기준(공통 종목 ÷ 작은 쪽 종목 수)**으로 나눠 표시하고, 그 한계를 화면에 그대로 적는다.",
        "**생성형 AI가 '아는 것만' 말한다** — AI에는 화면이 계산한 근거(구성 비율·겹침·발견)만 넘기고, 돌아온 문장 속 숫자를 근거 숫자와 대조해 근거에 없는 숫자가 든 문장을 자동 삭제한다(사후 대조). 매수·매도·추천·전망은 프롬프트로 금지한다. AI 서버가 없거나 실패해도 같은 근거로 규칙 진단서가 나가므로 서비스는 항상 답한다.",
        f"**이름 해독** — 상품명을 낱말 단위로 풀어 '무슨 뜻'과 '왜 신경 써야 하는지'를 붙인다(운용사 브랜드·상품 구조·지수·자산·지역·테마 {N_TERMS}낱말). 'TR = 분배금 재투자, 내용물은 같음', '커버드콜 = 분배금은 높지만 상승이 제한됨'처럼 매수 화면이 말해 주지 않는 것을 말한다.",
        "**팔지 않는 비교** — '같은 지수 다른 ETF'는 순자산·분배수익률·1년 수익률·변동성만 나란히 놓고 추천하지 않는다. 금융소비자보호법의 부당권유 우려 없이 증권사·은행이 그대로 붙일 수 있는 구조다.",
        "**계산은 코드가, 설명은 AI가** — 겹침률·실효 비중은 계산 문제라 AI에 맡기면 틀리고 재현되지 않는다. 계산·판정은 규칙 코드가 결정적으로 하고 AI는 설명만 맡는 분업으로 금융 서비스의 신뢰를 지킨다.",
        "**재사용 가능한 지식베이스** — 미래에셋 AI 페스티벌(2026.8~9)에서 팀으로 구축한 금융상품 지식베이스(온톨로지·구성종목 그래프·해석 사전)를 '질문에 답하는 API'가 아니라 '보유 바구니를 진단하는 서비스'라는 다른 용도로 활용했다. 화면·계산·발견 규칙·AI 진단 서버·이름 해독 사전은 이 대회를 위해 새로 만들었다.",
    ]:
        rich(c, s, bullet=True)

    c = add_section(t, "5. 활용 데이터 및 생성형 AI 모델 적용 방안")
    rich(c, "**활용 데이터와 수집·정제 방안**", 10.5)
    sub_table(c, ["데이터", "내용", "출처·수집 방법", "규모"], [
        ["① 국내 상장 ETF 마스터", "이름·운용사·기초지수·자산유형·투자지역·위험등급(1~6)·순자산·분배수익률·분배주기·1년 수익률·1년 변동성·복제방식·배수·상장일", "미래에셋증권 AI Festival 제공 금융상품 데이터(주최 재배포본, 기준일 2026-08-22). 상장 중인 ETF만 사용(ETN·상장폐지 제외)", f"{meta['etfs']:,}종"],
        ["② ETF 구성종목", "ETF별 편입 종목·비중(국내 주식)·종목 이름(해외 주식)", "한국거래소 정보데이터시스템 구성종목(PDF) 공시를 종목별로 직접 수집(기준일 2026-08-21). 현금·예금·환헤지 포워드 행 제외, ISIN으로 종목 통합(운용사마다 다른 표기 통일)", f"{meta['holding_pairs']:,}건 · 종목 {meta['securities']:,}개(비중 있음 {meta['pairs_with_weight']:,}건)"],
        ["③ 이름 해독 사전", "운용사 브랜드·상품 구조·지수·자산·지역·테마 낱말의 뜻과 주의점", "직접 작성(운용사 공시·거래소 안내·금융소비자보호법 위험등급 체계 참고) + 해외 종목 한글 별칭 82건", f"{N_TERMS}낱말"],
    ], [3.0, 5.2, 5.6, 2.6])
    for s in [
        "**정제 원칙** — 값이 없는 것은 채우지 않는다(총보수는 원천 값 보유율이 18%라 쓰지 않았다). 해외 주식 구성종목은 비중이 없어 '종목 이름'으로만 쓰고 화면에 그 사실을 적는다. 데이터는 build_data.py 한 번으로 재생성된다.",
        "**개인정보** — 수집하지 않는다. 바구니는 사용자의 브라우저 저장소와 주소(해시)에만 있고, AI 서버에는 상품명과 비율만 보낸다.",
    ]:
        rich(c, s, bullet=True)
    rich(c, "**생성형 AI 모델의 역할과 적용 방식**", 10.5)
    sub_table(c, ["단계", "무엇을 하나", "어떻게"], [
        ["1. 근거 만들기(코드)", "바구니 → 구성 비율·겹침 쌍·공통 종목·실효 비중·발견 목록", "브라우저 안 규칙 계산(결정적·재현 가능). AI는 계산에 관여하지 않음"],
        ["2. 진단서 쓰기(생성형 AI)", "근거를 쉬운 말 3단 진단서로: [한눈에 보기] · [발견한 점] · [확인해 볼 질문]", "**HyperCLOVA X(HCX-005)**, 온도 0.3·seed 고정. 시스템 프롬프트: 근거 밖 숫자·상품명 금지, 매수·매도·추천·전망 금지, 전문용어 괄호 풀이, 형식·면책 문구 고정. 서버 설정만으로 Claude·OpenAI 호환 모델로 교체 가능"],
        ["3. 사후 대조(코드)", "AI 문장의 숫자를 근거 숫자 집합과 대조", "근거에 없는 숫자가 든 줄은 삭제하고 삭제 건수를 화면에 표시. 남은 글이 80자 미만이면 규칙 진단서로 교체"],
        ["4. 항상 답하기", "AI 키 없음·호출 실패·40초 초과·분당 12회 초과", "같은 근거로 규칙 진단서(사람이 정한 문장 틀)를 돌려줌 — 서비스는 멈추지 않음"],
    ], [3.4, 5.0, 8.0])
    rich(c, "**확장 역할(설계 완료, MVP 범위 밖)** — 사전에 없는 새 낱말의 해설 초안 생성, '이 두 개 중 뭐가 겹쳐?' 같은 대화형 질문은 미래에셋 AI 페스티벌에서 만든 금융상품 질의응답 엔진(근거 기반, 답변 불가 시 거절)과 연결한다.", bullet=True)

    c = add_section(t, "6. 기대 효과 및 확장 가능성")
    rich(c, "**해결하는 문제와 기대 효과**", 10.5)
    for s in [
        "**고객(청년 투자자)** — '가짜 분산'과 한 종목 집중을 스스로 발견해 불필요한 중복 매수를 줄이고 의도한 분산으로 바꾼다. 상품명을 이해하고 사게 되어 커버드콜·레버리지 오해로 인한 손실을 예방한다. 비용 0원·가입 0회로, 소액 투자자도 PB급 '포트폴리오 엑스레이'를 받는다.",
        "**서비스 제공자(증권사·은행·핀테크)** — 상담 인력 없이 청년 고객에게 '보유상품 진단' 경험을 제공해 앱 체류·재방문을 높인다. 추천이 아닌 사실 설명 구조라 부당권유 규제 우려가 낮고, 상담원이 근거 표를 그대로 인용하는 보조 도구로도 쓸 수 있다.",
        "**사회** — 청년 자산형성을 돕는 포용금융. 예시 바구니와 이름 해독은 그대로 금융교육 자료가 된다.",
    ]:
        rich(c, s, bullet=True)
    rich(c, "**시장 확대·추가 기능·서비스 확장**", 10.5)
    for s in [
        "**상품군 확장** — 미래에셋 AI 페스티벌 지식베이스에 이미 적재된 공모펀드 23,622종·해외 상장 ETF 6,037종·국내채권 20,497종으로 확대. 펀드 클래스와 해외 ETF 구성종목(미국 SEC 공시) 연결 시 '계좌 전체 투시'가 된다.",
        "**데이터 확장** — 마이데이터 연동으로 바구니 자동 입력, 매월 기준일 갱신 자동화(구성종목 수집기 보유), 총보수·괴리율·추적오차 추가.",
        "**기능 확장** — 기준일별 겹침 변화(시간축), 목적별 점검(비상금·주택·노후 질문 세트), 친구와 바구니 비교, 상담원 모드(근거 표 인쇄), 고령층을 위한 큰 글씨·음성 읽기 모드.",
    ]:
        rich(c, s, bullet=True)
    rich(c, "**금융 서비스 외 타 영역 응용**", 10.5)
    rich(c, "'계산은 규칙, 설명은 AI, 숫자는 사후 대조'라는 구조는 그대로 옮겨 쓸 수 있다 — 보험 보장 겹침(중복 가입 진단), 카드 혜택 겹침, 통신·구독 요금제 중복, 대출 상품 조건 설명. 어느 영역이든 '이미 가진 것들의 관계'를 투시해 주는 소비자 보호형 AI 서비스가 된다.", bullet=True)

    c = add_section(t, "7. MVP 제작 과정과 정직한 한계", required=False)
    for s in [
        "**제작 과정** — 미래에셋 AI 페스티벌(2026.8~9)에서 3인 팀으로 구축한 금융상품 지식베이스(온톨로지 5파일·지식그래프 92만 트리플·한국거래소 구성종목 수집기·해석 사전)를 바탕으로, 본 대회에서는 1인이 '보유 바구니 진단'이라는 다른 용도의 서비스를 새로 설계·구현했다. 화면(모바일 우선 단일 페이지), 투시 계산, 발견 규칙 15종, AI 진단 서버(근거 → HyperCLOVA X → 사후 대조 → 규칙 진단서 대체), 이름 해독 사전은 모두 본 대회를 위해 새로 작성했다.",
        "**실측** — KODEX 200 ↔ TIGER 200 겹침 99.8%(비중 기준), KODEX 미국나스닥100 ↔ TIGER 미국나스닥100 100%(종목 수 기준), 국내 대표지수 중복형 예시 바구니에서 삼성전자 실효 비중 30.8%·SK하이닉스 25.0%, HyperCLOVA X 진단서 작성 약 9초·근거 대조로 지운 문장 0개(2026-09-07 실측).",
        "**정직한 한계** — 해외 주식 구성종목은 비중이 없어 종목 수 기준 겹침만 계산한다. 총보수는 표시하지 않는다. 기준일이 2026-08-22로 고정되어 그 뒤의 상장·가격 변화는 반영되지 않는다. 액티브 ETF는 구성이 자주 바뀐다. AI 진단서는 같은 근거라도 표현이 조금씩 다를 수 있다(온도 0.3). 이 서비스는 투자 권유가 아니며 판단은 사용자의 몫이다.",
        f"**주소** — 웹서비스 {URL} · 소스코드 {REPO}",
    ]:
        rich(c, s, bullet=True)
    out = os.path.join(DOCS, "기획서_김주연등장.docx")
    doc.save(out)
    return out


# ---------------- 첨부 2 기능명세서 ----------------
def build_spec():
    doc = new_doc()
    header_block(doc, 2, "2026 금융 AI Challenge 기능 명세서")
    t = section_table(doc)

    c = add_section(t, "1. MVP 구현 범위")
    rich(c, f"**서비스** ETF 엑스레이(ETF X-ray) · **배포 URL** {URL} · **소스코드** {REPO}", 9.5)
    for s in [
        f"**대상 데이터** — 국내 상장 ETF {meta['etfs']:,}종(상장 중, 기준일 2026-08-22)과 그 구성종목 {meta['holding_pairs']:,}건(기준일 2026-08-21). 모두 정적 파일(JSON)로 화면과 함께 배포되어 브라우저 안에서 계산한다.",
        "**구현·동작하는 범위** — ① ETF 검색·바구니 담기·금액 입력·예시 바구니 4종 ② 투시 계산(구성 비율 3종, 가중 지표 4종, ETF 쌍 겹침, 공통 종목, 실효 비중) ③ 발견 규칙 15종 ④ 생성형 AI 진단서(HyperCLOVA X + 숫자 사후 대조 + 규칙 진단서 대체) ⑤ 이름 해독(사전 " + str(N_TERMS) + "낱말) ⑥ 같은 지수 다른 ETF 비교표 ⑦ 바구니 주소 공유·브라우저 저장 ⑧ AI 서버 상태 확인.",
        "**구성** — 정적 화면(GitHub Pages, index.html + assets + data) 과 AI 진단서 서버(FastAPI, api/main.py). 화면은 서버 없이도 모든 투시 기능이 동작하고, AI 진단서만 서버가 있을 때 생성형 AI로, 없을 때 규칙 진단서로 나간다." + (f" {AI_NOTE}" if AI_NOTE else ""),
    ]:
        rich(c, s, bullet=True)

    c = add_section(t, "2. 주요 기능 목록")
    sub_table(c, ["번호", "기능명", "기능 설명", "관련 화면", "구현 상태"], [
        ["F1", "ETF 검색·담기", "이름·브랜드 별칭(타이거→TIGER, 코덱스→KODEX 등)·종목코드·기초지수로 검색, 순자산 순 8개 제안, 담기·금액(만원) 입력·빼기·비우기. Enter 키로 첫 제안 담기", "①바구니", "완료"],
        ["F2", "예시 바구니", "미국 빅테크 몰빵형·국내 대표지수 중복형·월배당 커버드콜형·골고루 균형형 4종을 한 번에 담고 바로 투시", "①바구니", "완료"],
        ["F3", "한눈에 보기", "ETF 수·서로 다른 종목 수·2개 이상 ETF에 겹친 종목 수·위험등급 가중평균·분배수익률·1년 변동성(가중) 6개 지표와 투자 지역·자산 유형·위험등급 비율 막대", "②한눈에", "완료"],
        ["F4", "ETF 쌍 겹침", "모든 ETF 쌍의 겹침률. 둘 다 비중 공개면 비중 기준(종목별 비중 중 작은 값의 합), 아니면 종목 수 기준(공통 종목 ÷ 작은 쪽 종목 수). '같은 지수'·'사실상 같은 상품(70% 이상)' 표시", "③겹침", "완료"],
        ["F5", "공통 종목·실효 비중", "2개 이상 ETF에 든 종목 12개(담은 ETF 수·담은 ETF 금액 비중·확인된 실효 비중)와 실효 비중 TOP10(비중 공개 ETF 기준, 금액 비중 표시)", "③겹침", "완료"],
        ["F6", "발견 규칙", "같은 지수 중복 · 70% 이상 겹침 · 포함 관계(작은 ETF가 큰 ETF 안에) · 40~70% 겹침 · 공통 종목 · 한 종목 집중 · 지역 몰림 · 전부 주식형 · 레버리지/인버스 · 위험 1등급 비중 · 커버드콜 · 합성복제 · 소규모 순자산 · 고변동성 · 월배당 (+ 이상 없음, ETF 1개 안내)", "④발견", "완료"],
        ["F7", "AI 진단서", "근거 → HyperCLOVA X(HCX-005) → 숫자 사후 대조 → [한눈에 보기]·[발견한 점]·[확인해 볼 질문] 3단 진단서. 작성 모델·소요 시간·지운 문장 수 표시, 'AI에 넘긴 근거 보기'. 서버 없음·실패·시간 초과 시 규칙 진단서", "⑤AI 진단서", "완료"],
        ["F8", "이름 해독", f"상품명을 사전 {N_TERMS}낱말과 대조해 뜻·왜 중요한지 표시. 상품 기본 정보(운용사·기초지수·복제방식·위험등급·분배·순자산·상장일) 한 줄", "⑥이름 해독", "완료"],
        ["F9", "같은 지수 다른 ETF", "기초지수가 같은 상장 ETF 최대 8개를 순자산·분배수익률·1년 수익률·1년 변동성·위험등급으로 비교(추천 없음)", "⑦같은 지수", "완료"],
        ["F10", "공유·저장", "바구니를 주소 해시(#b=종목코드:금액)에 담아 복사, 브라우저 저장소에서 복원. 주소만으로 같은 결과 재현", "⑧공유", "완료"],
        ["F11", "AI 서버 상태", "/api/health 로 서버·모델 확인, 연결 시 화면 하단에 표시", "하단", "완료"],
    ], [1.0, 2.6, 8.4, 2.4, 1.8], size=8)
    rich(c, "**화면 예시(모바일)**", 9.5)
    images_row(c, [os.path.join(IMG, f) for f in ("shot1_basket.png", "shot2_overview.png", "shot3_overlap.png")],
               ["① 바구니 담기·예시 바구니", "② 한눈에 보기(구성 비율)", "③ 겹침 투시(쌍 겹침·공통 종목)"], 5.4)
    images_row(c, [os.path.join(IMG, f) for f in ("shot4_findings.png", "shot5_ai.png", "shot6_terms.png")],
               ["④ 발견한 점", "⑤ AI 진단서(HyperCLOVA X)", "⑥ 이름 해독·같은 지수"], 5.4)

    c = add_section(t, "3. 사용자 이용 흐름")
    for i, s in enumerate([
        f"배포 URL({URL}) 접속 → 첫 화면 '내 ETF 바구니 담기'. 로그인·설치 없음.",
        "예시 바구니 칩(예: 🇰🇷 국내 대표지수 중복형)을 누르면 ETF 5개가 담기고 바로 투시된다. 직접 담으려면 검색창에 '타이거 나스닥'처럼 입력 → 제안 목록에서 선택 → 금액(만원) 입력.",
        "'투시하기' → 결과 6개 섹션이 열리고 상단에 탭(한눈에·겹침·발견·AI 진단서·이름 해독·같은 지수)이 고정된다.",
        "'한눈에'에서 지표 6개와 투자 지역·자산·위험등급 비율을 본다 → '겹침'에서 ETF 쌍 겹침률과 여러 ETF에 함께 든 종목·실효 비중을 본다 → '발견'에서 주의(주황)·참고(파랑) 카드를 읽는다.",
        "'AI 진단서 받기' → 10초 안팎 뒤 생성형 AI가 쓴 진단서가 표시된다(작성 모델·소요 시간·지운 문장 수 표기). 'AI에 넘긴 근거 보기'로 입력 근거를 확인할 수 있다.",
        "'이름 해독'에서 각 ETF 이름 속 낱말의 뜻을, '같은 지수'에서 같은 지수를 따르는 다른 ETF 비교표를 본다.",
        "'공유'의 주소를 복사해 친구에게 보낸다. 받은 사람은 주소만 열면 같은 바구니가 바로 투시된다. 바구니를 고치고 다시 '투시하기'를 누르면 갱신된다.",
    ], 1):
        rich(c, f"{i}) {s}", indent=0.4)

    c = add_section(t, "4. AI 및 데이터 처리 방식")
    rich(c, "**AI가 수행하는 역할**", 10.5)
    for s in [
        "생성형 AI는 **설명만** 맡는다. 화면(코드)이 계산한 근거를 받아 사회초년생이 이해할 수 있는 3단 진단서를 쓴다. 겹침률·실효 비중·발견 판정 같은 계산은 AI가 하지 않는다(결정적 규칙 코드).",
        "모델: 네이버 **HyperCLOVA X HCX-005**(CLOVA Studio v3 chat-completions), 온도 0.3·topP 0.8·seed 7·최대 1,200토큰. 서버 환경변수만 바꾸면 Claude·OpenAI 호환 모델로 교체된다.",
        "시스템 프롬프트 규칙: ① 근거에 있는 사실·숫자만 사용 ② 매수·매도·추천·전망 금지 ③ 전문용어는 괄호로 풀이 ④ [한눈에 보기]·[발견한 점]·[확인해 볼 질문] 형식과 면책 문구 고정.",
        "사후 대조: 응답의 모든 숫자를 근거 숫자 집합과 비교해 근거에 없는 숫자가 든 줄을 삭제하고 건수를 화면에 표시한다. 남은 글이 80자 미만이면 규칙 진단서로 교체한다. 호출 실패·40초 초과·키 없음·IP당 분당 12회 초과 시에도 규칙 진단서로 답한다.",
    ]:
        rich(c, s, bullet=True)
    rich(c, "**사용 데이터와 입력·출력**", 10.5)
    sub_table(c, ["구분", "내용"], [
        ["화면 데이터", f"data/etfs.json — 상장 ETF {meta['etfs']:,}종 속성(약 0.6MB) · data/holdings.json — 구성종목 {meta['holding_pairs']:,}건 + 종목 사전 {meta['securities']:,}개(약 1.3MB). 첫 접속 때 한 번 내려받아 브라우저 안에서 계산"],
        ["AI 입력", "바구니(이름·투자 비중·운용사·기초지수·위험등급·분배수익률) · 총 투자금액 · 구성 비율(지역·자산·위험등급) · 여러 ETF에 든 종목 8개(담은 ETF 수·실효 비중) · 겹침 쌍 6개 · 비중 확인 금액 비중 · 발견 목록 — 모두 줄글로 만들어 [근거]로 전달"],
        ["AI 출력", "500자 안팎 진단서 3단 + 면책 문구 → 사후 대조 → 화면 표시(작성 모델·소요 시간·지운 문장 수)"],
        ["데이터 출처", "국내 상장 ETF 마스터: 미래에셋증권 AI Festival 제공(기준일 2026-08-22) · 구성종목: 한국거래소 정보데이터시스템에서 직접 수집(기준일 2026-08-21) · 이름 해독 사전: 직접 작성"],
        ["개인정보·민감정보", "수집·저장·전송하지 않는다. 로그인 없음. 바구니는 사용자 브라우저 저장소와 주소 해시에만 있으며, AI 서버에는 상품명·비율만 전송된다(식별 정보 없음). 서버는 요청 내용을 저장하지 않는다."],
    ], [3.2, 13.2])

    c = add_section(t, "5. MVP 검증 방법")
    rich(c, "**심사자 확인 절차(약 3분, 계정 불필요)**", 10.5)
    steps = [
        f"{URL} 접속 → 첫 화면에 '내 ETF 바구니 담기'와 예시 바구니 칩 4개가 보인다.",
        "'🇰🇷 국내 대표지수 중복형' 클릭 → 바구니에 ETF 5개(합계 1,000만원)가 담기고 결과가 열린다. 예상: 한눈에 보기 'ETF 개수 5', 겹침 첫 줄 'KODEX 200 ↔ KODEX 레버리지 100% 종목 수 기준·같은 지수', 'KODEX 200 ↔ TIGER 200 99.8% 비중 기준', 공통 종목 1위 '삼성전자 5개 100% 30.8%', 발견 첫 카드 '같은 지수를 따르는 ETF를 3개 갖고 있어요'.",
        "'AI 진단서 받기' 클릭 → 10초 안팎 뒤 진단서 표시. 예상: 상태줄 '✅ HyperCLOVA X가 작성 · N초 · 근거 대조로 지운 문장 N개', 본문에 [한눈에 보기]·[발견한 점]·[확인해 볼 질문]과 면책 문구. AI 서버가 연결되지 않은 환경이면 'ℹ️ AI 서버에 연결되지 않아 규칙 기반 진단서로 대신했어요'와 함께 같은 형식의 진단서가 나온다.",
        "검색창에 '타이거 나스닥' 입력 → 제안 'TIGER 미국나스닥100' 선택 → 금액 200 입력 → '투시하기'. 예상: 투자 지역 막대에 '미국'이 생기고, 겹침 표에 TIGER 미국나스닥100과 다른 ETF의 쌍(종목 수 기준)이 추가된다.",
        "'🇺🇸 미국 빅테크 몰빵형' 클릭 → 예상: 'TIGER 미국나스닥100 ↔ KODEX 미국나스닥100 100% 종목 수 기준·같은 지수', 공통 종목 1위 '엔비디아 (NVIDIA CORP) 4개 100% 비중 비공개', 발견에 '투자 지역의 100%가 미국에 몰려 있어요'. '⚖️ 골고루 균형형' 클릭 → 예상: 겹친 종목 0, 발견 '크게 겹치거나 한쪽에 몰린 부분은 보이지 않아요'.",
        "'공유' 주소 복사 → 새 창(또는 다른 기기)에서 열기 → 같은 바구니가 바로 투시된다.",
    ]
    for i, s in enumerate(steps, 1):
        rich(c, f"{i}) {s}", indent=0.4)
    if SAMPLES:
        rich(c, "**예시 바구니별 예상 결과(2026-09-07 실측)**", 10.5)
        rows = []
        for name, r in SAMPLES.items():
            rows.append([name, r.get("etfs", ""), r.get("kpi", ""), r.get("top_pair", ""), r.get("top_common", ""), r.get("first_finding", "")])
        sub_table(c, ["예시 바구니", "ETF", "지표", "겹침 1위", "공통 종목 1위", "발견 첫 카드"], rows, [2.6, 3.6, 2.6, 2.8, 2.4, 2.4], size=7.5)
    rich(c, "**테스트 계정·샘플 입력값** — 계정 불필요. 샘플 입력: 예시 바구니 4종(위 표) 또는 검색어 '타이거 나스닥', 'KODEX 200', '069500', '월배당', '커버드콜'. 주소로 바로 열기: " + URL + "#b=069500:300,102110:300,278530:200,122630:100,091160:100 (국내 대표지수 중복형과 같은 바구니) — 끝에 &ai=1 을 붙이면 AI 진단서까지 자동 요청한다.", bullet=True)
    rich(c, "**실행 환경·브라우저** — 최신 Chrome·Edge·Safari·Samsung Internet(모바일 포함). JavaScript 필요. 첫 접속 때 약 2MB를 내려받는다(이후 브라우저 캐시). 별도 설치·확장 프로그램 없음.", bullet=True)
    rich(c, "**MVP 단계의 제한사항** — ① 대상은 국내 상장 ETF만이며 기준일이 2026-08-22로 고정되어 그 뒤의 상장·가격 변화는 없다. ② 해외 주식 구성종목은 원천에 비중이 없어 '종목 수 기준' 겹침만 계산하고 실효 비중에서 빠진다(화면에 표기). ③ 총보수는 표시하지 않는다. ④ AI 진단서는 서버 상태·호출량(IP당 분당 12회)에 따라 규칙 기반 진단서로 대체될 수 있고, 같은 근거라도 표현이 조금씩 다를 수 있다. ⑤ 이 서비스는 투자 권유가 아니다." + (f" ⑥ {AI_NOTE}" if AI_NOTE else ""), bullet=True)
    out = os.path.join(DOCS, "기능명세서_김주연등장.docx")
    doc.save(out)
    return out


def to_pdf(docx_path):
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        pdf = os.path.splitext(os.path.abspath(docx_path))[0] + ".pdf"
        d.ExportAsFixedFormat(pdf, 17)   # 17 = wdExportFormatPDF
        d.Close(False)
        return pdf
    finally:
        word.Quit()


if __name__ == "__main__":
    outs = [build_plan(), build_spec()]
    for o in outs:
        print("docx:", o)
        if "--no-pdf" not in sys.argv:
            print("pdf :", to_pdf(o))
